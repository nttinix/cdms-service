import io
import json
import zipfile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from cdms.models import Product

COLUMNS = [
    "productId",
    "sourceVersion",
    "sku",
    "partnerSKU",
    "productName",
    "description",
    "color",
    "size",
    "isActive",
    "units",
    "categories",
]
REQUIRED = {"productId", "sourceVersion", "sku", "partnerSKU", "productName"}


class SpreadsheetError(ValueError):
    def __init__(self, message, row=None):
        self.row = row
        super().__init__(message)


def decode_workbook(content, max_rows=10_000):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise SpreadsheetError("Expanded workbook exceeds 100 MiB")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (
        zipfile.BadZipFile,
        InvalidFileException,
        OSError,
        ValueError,
        KeyError,
        ParseError,
    ) as error:
        raise SpreadsheetError("Invalid XLSX workbook") from error
    try:
        sheet = book.active
        rows = sheet.iter_rows()
        header_cells = next(rows, None)
        headers = tuple(cell.value for cell in header_cells) if header_cells else None
        if headers is None or any(not isinstance(value, str) for value in headers):
            raise SpreadsheetError("A header row is required", 1)
        if len(set(headers)) != len(headers) or not REQUIRED.issubset(headers):
            raise SpreadsheetError("Missing required or duplicated columns", 1)
        if not set(headers).issubset(COLUMNS):
            raise SpreadsheetError("Unknown columns", 1)
        products, row_numbers = [], []
        for row_number, cells in enumerate(rows, start=2):
            if row_number > max_rows + 1:
                raise SpreadsheetError("Workbook exceeds row limit", row_number)
            if any(cell.data_type == "f" for cell in cells):
                raise SpreadsheetError("Formula cells are not accepted", row_number)
            values = tuple(cell.value for cell in cells)
            if all(value is None for value in values):
                continue
            data = dict(zip(headers, values, strict=True))
            try:
                for column in ("productId", "sourceVersion"):
                    value = data[column]
                    if isinstance(value, (int, float)) and abs(value) >= 10**15:
                        raise ValueError("Large Excel integers must be stored as text")
                    if isinstance(value, str) and 16 <= len(value) <= 19:
                        if value.isascii() and value.isdigit() and value == str(int(value)):
                            data[column] = int(value)
                for column in ("units", "categories"):
                    if column in data:
                        data[column] = json.loads(data[column]) if data[column] else []
                products.append(Product.model_validate(data))
                row_numbers.append(row_number)
            except (ValidationError, ValueError, TypeError) as error:
                raise SpreadsheetError("Invalid product data", row_number) from error
        if not products:
            raise SpreadsheetError("Workbook contains no products")
        return products, row_numbers
    except (ParseError, zipfile.BadZipFile, OSError, KeyError) as error:
        raise SpreadsheetError("Invalid worksheet data") from error
    finally:
        book.close()


def encode_workbook(products):
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Products")
    sheet.append(COLUMNS)
    for product in products:
        row = []
        for column in COLUMNS:
            value = product.get(column)
            if column in {"productId", "sourceVersion"} and type(value) is int and value >= 10**15:
                value = str(value)
            if isinstance(value, (list, dict)):
                value = json.dumps(value, ensure_ascii=False)
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                from openpyxl.cell import WriteOnlyCell

                cell = WriteOnlyCell(sheet, value=value)
                cell.data_type = "s"
                value = cell
            row.append(value)
        sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()
